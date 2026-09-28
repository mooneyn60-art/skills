"""Summarize KDP royalty/sales reports exported from KDP Reports (Prior Month Royalties,
Month-to-Date, or Dashboard downloads) as .xlsx or .csv.

Column names are matched loosely, so it works across KDP's report variants:
title, marketplace, royalty type/format, units (net units sold), royalty, currency, KENP read.

Usage:
  python royalty_report.py KDP_Prior_Month_Royalties-2026-08.xlsx
  python royalty_report.py report1.xlsx report2.xlsx --xlsx summary.xlsx
"""

import argparse
import csv
import json
import os
import re
from collections import defaultdict

FIELDS = {
    "title": [r"^title$"],
    "marketplace": [r"marketplace"],
    "format": [r"royalty type", r"^format$", r"transaction type"],
    "units": [r"net units sold", r"^units sold$", r"net units"],
    "royalty": [r"^royalty$", r"royalty amount", r"^royalties$"],
    "currency": [r"currency"],
    "kenp": [r"kenp", r"normalized page"],
    "date": [r"date"],
}


def map_columns(header):
    names = [str(h or "").strip().lower() for h in header]
    mapping = {}
    for field, patterns in FIELDS.items():
        for pattern in patterns:
            idx = next((i for i, n in enumerate(names) if re.search(pattern, n)), None)
            if idx is not None:
                mapping[field] = idx
                break
    return mapping


def rows_from_file(path):
    """Yield (sheet_name, rows) with raw row lists."""
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            yield os.path.basename(path), list(csv.reader(f))
        return
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)]


def to_number(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^\d.\-]", "", str(value))
    return float(cleaned) if cleaned not in ("", "-", ".") else 0.0


def parse(paths):
    records, kenp = [], []
    for path in paths:
        for sheet, rows in rows_from_file(path):
            header_idx = next((i for i, r in enumerate(rows[:30])
                               if "title" in map_columns(r) and ({"royalty", "units", "kenp"} & map_columns(r).keys())), None)
            if header_idx is None:
                continue
            cols = map_columns(rows[header_idx])
            get = lambda r, k: r[cols[k]] if k in cols and cols[k] < len(r) else None
            for r in rows[header_idx + 1:]:
                title = get(r, "title")
                if not title or str(title).strip().lower() in ("total", "totals"):
                    continue
                base = {"title": str(title).strip(), "marketplace": str(get(r, "marketplace") or "unknown"),
                        "sheet": sheet}
                if "kenp" in cols and "royalty" not in cols:
                    kenp.append({**base, "pages": to_number(get(r, "kenp"))})
                else:
                    records.append({**base, "format": str(get(r, "format") or sheet),
                                    "units": to_number(get(r, "units")), "royalty": to_number(get(r, "royalty")),
                                    "currency": str(get(r, "currency") or "unknown")})
    return records, kenp


def summarize(records, kenp):
    by_title = defaultdict(lambda: {"units": 0.0, "royalty_by_currency": defaultdict(float), "kenp_pages": 0.0})
    by_market = defaultdict(lambda: defaultdict(float))
    by_format = defaultdict(lambda: {"units": 0.0, "royalty_by_currency": defaultdict(float)})
    totals = defaultdict(float)
    for r in records:
        t = by_title[r["title"]]
        t["units"] += r["units"]
        t["royalty_by_currency"][r["currency"]] += r["royalty"]
        by_market[r["marketplace"]][r["currency"]] += r["royalty"]
        f = by_format[r["format"]]
        f["units"] += r["units"]
        f["royalty_by_currency"][r["currency"]] += r["royalty"]
        totals[r["currency"]] += r["royalty"]
    for k in kenp:
        by_title[k["title"]]["kenp_pages"] += k["pages"]

    def clean(d):
        return {k: round(v, 2) for k, v in d.items()}

    titles = sorted(by_title.items(), key=lambda kv: -sum(kv[1]["royalty_by_currency"].values()))
    return {
        "royalty_totals_by_currency": clean(totals),
        "kenp_pages_total": round(sum(k["pages"] for k in kenp)),
        "titles": [{"title": name, "units": round(v["units"]), "royalty_by_currency": clean(v["royalty_by_currency"]),
                    "kenp_pages": round(v["kenp_pages"])} for name, v in titles],
        "marketplaces": {m: clean(c) for m, c in sorted(by_market.items())},
        "formats": {f: {"units": round(v["units"]), "royalty_by_currency": clean(v["royalty_by_currency"])}
                    for f, v in sorted(by_format.items())},
        "rows_parsed": len(records) + len(kenp),
        "note": "Royalties are not converted between currencies. KENP payout rate is set monthly by Amazon.",
    }


def write_xlsx(summary, path):
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "By title"
    currencies = sorted(summary["royalty_totals_by_currency"])
    ws.append(["Title", "Units", "KENP pages"] + [f"Royalty {c}" for c in currencies])
    for t in summary["titles"]:
        ws.append([t["title"], t["units"], t["kenp_pages"]] + [t["royalty_by_currency"].get(c, 0) for c in currencies])
    m = wb.create_sheet("By marketplace")
    m.append(["Marketplace"] + [f"Royalty {c}" for c in currencies])
    for name, c in summary["marketplaces"].items():
        m.append([name] + [c.get(x, 0) for x in currencies])
    f = wb.create_sheet("By format")
    f.append(["Format", "Units"] + [f"Royalty {c}" for c in currencies])
    for name, v in summary["formats"].items():
        f.append([name, v["units"]] + [v["royalty_by_currency"].get(x, 0) for x in currencies])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 45
    wb.save(path)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("reports", nargs="+")
    p.add_argument("--xlsx", help="also write a summary workbook")
    a = p.parse_args()
    records, kenp = parse(a.reports)
    if not records and not kenp:
        raise SystemExit("No KDP report tables found. Export from KDP Reports > Prior Month Royalties or Dashboard.")
    summary = summarize(records, kenp)
    if a.xlsx:
        write_xlsx(summary, a.xlsx)
        summary["xlsx"] = a.xlsx
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
